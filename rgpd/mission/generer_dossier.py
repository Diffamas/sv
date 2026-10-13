#!/usr/bin/env python3
"""Génère le dossier RGPD complet d'un client à partir de son fichier YAML.

    python3 generer_dossier.py client.yaml [--sans-audit]

Produit dossiers/<client>/ :
  00 à 07 : les documents du kit, remplis avec les données du client
  rapport-audit-site.md : audit express du site (si `site` est renseigné)
  A-COMPLETER.txt : la liste des champs restés entre [crochets], fichier par fichier
"""
import argparse
import datetime as dt
import os
import re
import shutil
import subprocess
import sys

import yaml
from docx import Document
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
KIT = os.path.join(HERE, "..", "kit-rgpd-tpe", "generer_kit.py")
AUDIT = os.path.join(HERE, "..", "outil-audit")


def slug(s):
    s = re.sub(r"[^a-z0-9]+", "-", s.lower().strip())
    return s.strip("-") or "client"


def build_map(c):
    today = dt.date.today().strftime("%d/%m/%Y")
    g = lambda k, d="": str(c.get(k) or d)
    pairs = [
        ("[Raison sociale], [forme juridique], [adresse], SIRET [numéro], représentée par [prénom nom]",
         f"{g('raison_sociale')}, {g('forme_juridique')}, {g('adresse')}, SIRET {g('siret')}, représentée par {g('representant')}"),
        ("[Raison sociale], [adresse], SIRET [numéro]", f"{g('raison_sociale')}, {g('adresse')}, SIRET {g('siret')}"),
        ("[prénom nom, fonction]", f"{g('representant')}, {g('fonction')}"),
        ("[prénom nom, email]", f"{g('representant')}, {g('email')}"),
        ("[nom, téléphone, email]", g("prestataire_it")),
        ("[nom, email]", f"{g('representant')}, {g('email')}"),
        ("[Nom de l'hébergeur], [adresse], [téléphone]", g("hebergeur_mentions", g("hebergeur"))),
        ("[raison sociale ou nom et prénom pour une entreprise individuelle]", g("raison_sociale")),
        ("[SARL au capital de X € / entreprise individuelle]", g("forme_capital", g("forme_juridique"))),
        ("[RCS ville + numéro, ou SIREN pour les entreprises individuelles ; numéro au répertoire des métiers pour les artisans]", g("immatriculation")),
        ("[numéro] ou « TVA non applicable, article 293 B du CGI »", g("tva")),
        ("[Hébergeur du site, ex. OVH, o2switch]", g("hebergeur")),
        ("[Messagerie, ex. Google Workspace, Microsoft 365]", g("messagerie")),
        ("[Expert-comptable]", g("comptable")),
        ("[Logiciel de facturation / caisse]", g("logiciel_facturation")),
        ("[Agence web / webmaster]", g("agence_web")),
        ("[Télésurveillance / installateur caméras]", g("videosurveillance_prestataire")),
        ("[Installateur / télésurveillance]", g("videosurveillance_prestataire")),
        ("[Google Analytics / Matomo]", g("analytics")),
        ("[Oui / Non si Matomo configuré en exemption CNIL]", g("analytics_consentement")),
        ("[YouTube, Google Maps]", g("embeds")),
        ("[Pixel Meta, LinkedIn]", g("pixels")),
        ("[Aucun transfert. / Certains prestataires, comme [outil], sont situés aux États-Unis : le transfert est encadré par les clauses contractuelles types de la Commission européenne et, le cas échéant, par leur certification au Data Privacy Framework.]", g("transfert_hors_ue")),
        ("[Raison sociale]", g("raison_sociale")), ("[raison sociale]", g("raison_sociale")),
        ("[forme juridique]", g("forme_juridique")), ("[adresse complète]", g("adresse")),
        ("[adresse postale]", g("adresse")), ("[adresse]", g("adresse")), ("[ville]", g("ville")),
        ("[prénom nom]", g("representant")), ("[gérant]", g("representant")),
        ("[email]", g("email")), ("[téléphone]", g("telephone")),
        ("[hébergeur]", g("hebergeur")), ("[messagerie]", g("messagerie")), ("[expert-comptable]", g("comptable")),
        ("[logiciel de facturation]", g("logiciel_facturation")), ("[agence web]", g("agence_web")),
        ("[logiciel paie]", g("comptable")), ("[logiciel]", g("logiciel_facturation")), ("[CMS]", g("agence_web")),
        ("[Hébergeur]", g("hebergeur")), ("[outil statistiques]", g("analytics")), ("[Outil emailing]", g("outil_emailing", "")),
        ("[date]", today),
    ]
    return [(k, v) for k, v in pairs if v]


def apply(text, mapping):
    for k, v in mapping:
        if k in text:
            text = text.replace(k, v)
    return text


def fix_docx(path, mapping, titre):
    d = Document(path)
    paras = list(d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                paras.extend(cell.paragraphs)
    for p in paras:
        txt = p.text
        if txt.startswith("KIT RGPD TPE"):
            p.runs[0].text = titre
            for r in p.runs[1:]:
                r.text = ""
            continue
        if txt.startswith("Modèle version 1.0"):
            for r in p.runs:
                r.text = ""
            continue
        if "[" in txt:
            new = apply(txt, mapping)
            if new != txt and p.runs:
                p.runs[0].text = new
                for r in p.runs[1:]:
                    r.text = ""
    d.save(path)
    return sorted(set(re.findall(r"\[[^\]]{2,80}\]", "\n".join(p.text for p in Document(path).paragraphs) +
                                 "\n".join(c.text for t in Document(path).tables for r in t.rows for c in r.cells))))


def fix_xlsx(path, mapping, c):
    wb = load_workbook(path)
    keep = set(c.get("traitements") or [])
    ws = wb["Registre"]
    if keep:
        for r in range(ws.max_row, 1, -1):
            ref = str(ws.cell(r, 1).value or "")
            if ref and ref.split("-")[0] not in keep:
                ws.delete_rows(r)
    if not c.get("videosurveillance", True) and "SEC" in keep:
        pass
    rest = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "[" in cell.value:
                    cell.value = apply(cell.value, mapping)
                    rest.update(re.findall(r"\[[^\]]{2,80}\]", cell.value))
    wb.save(path)
    return sorted(rest)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("yaml")
    ap.add_argument("--sans-audit", action="store_true")
    a = ap.parse_args()
    c = yaml.safe_load(open(a.yaml, encoding="utf-8"))
    out = os.path.join(HERE, "dossiers", slug(c["raison_sociale"]))
    os.makedirs(out, exist_ok=True)
    subprocess.run([sys.executable, KIT], cwd=out, check=True, capture_output=True)
    mapping = build_map(c)
    titre = f"DOSSIER RGPD — {c['raison_sociale'].upper()}"
    report = []
    for fn in sorted(os.listdir(out)):
        p = os.path.join(out, fn)
        if fn.endswith(".docx"):
            rest = fix_docx(p, mapping, titre)
        elif fn.endswith(".xlsx"):
            rest = fix_xlsx(p, mapping, c)
        else:
            continue
        report.append((fn, rest))
    if c.get("site") and not a.sans_audit:
        sys.path.insert(0, AUDIT)
        from audit_site import audit  # noqa: E402
        r = audit(c["site"])
        lines = [f"# Audit express du site {c['site']}", f"Date : {dt.date.today():%d/%m/%Y}", "",
                 f"Score : {r['score']}", ""]
        for k in ("https", "mentions_legales", "politique_confidentialite", "politique_cookies", "cgv", "cmp",
                  "traceurs", "formulaire_contact", "consentement_formulaire"):
            lines.append(f"- {k.replace('_', ' ')} : {r[k]}")
        lines += ["", "## Constat", r["accroche"]]
        open(os.path.join(out, "rapport-audit-site.md"), "w", encoding="utf-8").write("\n".join(lines))
    with open(os.path.join(out, "A-COMPLETER.txt"), "w", encoding="utf-8") as f:
        f.write(f"Dossier généré le {dt.date.today():%d/%m/%Y} pour {c['raison_sociale']}.\n"
                "Champs restés entre crochets, à compléter à la main :\n\n")
        for fn, rest in report:
            f.write(f"{fn} : {', '.join(rest) if rest else 'rien'}\n")
    print(f"Dossier généré : {out}")
    for fn, rest in report:
        print(f"  {fn}: {len(rest)} champ(s) à compléter")


if __name__ == "__main__":
    main()
