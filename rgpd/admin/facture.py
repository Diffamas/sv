#!/usr/bin/env python3
"""Facturation et livre des recettes d'une micro-entreprise.

  python3 facture.py nouvelle --client "Boulangerie Dupont" --adresse "12 rue de la Gare, 77130 Montereau" \
      --objet "Pack Site conforme" --montant 290 [--siret-client ...] [--detail "ligne 1;ligne 2"]
      -> factures/2026-001.docx (+ .pdf si LibreOffice est installé), ligne ajoutée à livre-recettes.csv

  python3 facture.py encaissee 2026-001 [--date 2026-10-20] [--mode virement]
      -> marque la facture payée dans le livre des recettes

  python3 facture.py urssaf [--trimestre 2026-T4]
      -> chiffre d'affaires encaissé à déclarer et cotisations estimées

  python3 facture.py relances
      -> factures échues non payées, avec le texte de relance prêt à envoyer
"""
import argparse
import csv
import datetime as dt
import os
import shutil
import subprocess

import yaml
from docx import Document
from docx.shared import Pt, Cm, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
CFG = os.path.join(HERE, "config-vendeur.yaml")
LIVRE = os.path.join(HERE, "livre-recettes.csv")
FACT = os.path.join(HERE, "factures")
COLS = ["numero", "date", "client", "objet", "montant", "echeance", "encaisse_le", "mode", "fichier"]


def cfg():
    return yaml.safe_load(open(CFG, encoding="utf-8"))


def lire():
    if not os.path.exists(LIVRE):
        return []
    with open(LIVRE, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def ecrire(rows):
    with open(LIVRE, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLS)
        w.writeheader()
        w.writerows(rows)


def prochain_numero(rows):
    annee = dt.date.today().year
    n = sum(1 for r in rows if r["numero"].startswith(str(annee))) + 1
    return f"{annee}-{n:03d}"


def eur(x):
    return f"{float(x):,.2f} €".replace(",", " ").replace(".", ",")


def nouvelle(a):
    v = cfg()
    rows = lire()
    num = prochain_numero(rows)
    today = dt.date.today()
    ech = today + dt.timedelta(days=int(v.get("delai_paiement_jours", 30)))
    os.makedirs(FACT, exist_ok=True)
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2)
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    p = d.add_paragraph()
    r = p.add_run(f"FACTURE N° {num}")
    r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = RGBColor(0x15, 0x23, 0x3B)
    d.add_paragraph(f"Date d'émission : {today:%d/%m/%Y}\nDate d'échéance : {ech:%d/%m/%Y}")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = (f"Émetteur\n{v['nom']}\n{v['adresse']}\nSIRET {v['siret']}\n{v['email']}\n{v['telephone']}")
    t.rows[0].cells[1].text = (f"Client\n{a.client}\n{a.adresse}" + (f"\nSIRET {a.siret_client}" if a.siret_client else ""))
    d.add_paragraph()
    tb = d.add_table(rows=1, cols=2)
    tb.style = "Light Grid Accent 1"
    tb.rows[0].cells[0].text = "Désignation"
    tb.rows[0].cells[1].text = "Montant"
    c = tb.add_row().cells
    c[0].text = a.objet + ("\n" + "\n".join("- " + l.strip() for l in a.detail.split(";")) if a.detail else "")
    c[1].text = eur(a.montant)
    c = tb.add_row().cells
    c[0].text = "Total à payer (TVA non applicable, art. 293 B du CGI)"
    c[1].text = eur(a.montant)
    d.add_paragraph()
    d.add_paragraph(f"Règlement par virement sous {v.get('delai_paiement_jours', 30)} jours.\nIBAN : {v['iban']}   BIC : {v['bic']}")
    d.add_paragraph("En cas de retard de paiement, pénalités au taux de trois fois le taux d'intérêt légal et indemnité "
                    "forfaitaire pour frais de recouvrement de 40 € (art. L441-10 du Code de commerce). "
                    "Pas d'escompte pour paiement anticipé.").runs[0].font.size = Pt(8)
    d.add_paragraph(f"{v['nom']} — entreprise individuelle — SIRET {v['siret']} — TVA non applicable, art. 293 B du CGI").runs[0].font.size = Pt(8)
    path = os.path.join(FACT, f"{num}.docx")
    d.save(path)
    pdf = ""
    if shutil.which("soffice"):
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", FACT, path],
                       capture_output=True, timeout=120)
        if os.path.exists(path[:-5] + ".pdf"):
            pdf = path[:-5] + ".pdf"
    rows.append({"numero": num, "date": today.isoformat(), "client": a.client, "objet": a.objet,
                 "montant": f"{float(a.montant):.2f}", "echeance": ech.isoformat(), "encaisse_le": "", "mode": "",
                 "fichier": os.path.basename(pdf or path)})
    ecrire(rows)
    print(f"Facture {num} créée : {pdf or path}\nÉchéance : {ech:%d/%m/%Y}")


def encaissee(a):
    rows = lire()
    for r in rows:
        if r["numero"] == a.numero:
            r["encaisse_le"] = a.date or dt.date.today().isoformat()
            r["mode"] = a.mode
            ecrire(rows)
            print(f"Facture {a.numero} marquée encaissée le {r['encaisse_le']} ({a.mode}).")
            return
    print("Numéro introuvable.")


def trimestre_de(d):
    return f"{d.year}-T{(d.month - 1) // 3 + 1}"


def urssaf(a):
    v = cfg()
    rows = lire()
    tri = a.trimestre or trimestre_de(dt.date.today())
    ca = sum(float(r["montant"]) for r in rows if r["encaisse_le"] and trimestre_de(dt.date.fromisoformat(r["encaisse_le"])) == tri)
    taux = float(v.get("taux_cotisations", 25.6))
    if v.get("acre"):
        taux *= 0.75
    print(f"Trimestre {tri}\nChiffre d'affaires encaissé à déclarer : {eur(ca)}\n"
          f"Cotisations estimées ({taux:.1f} %) : {eur(ca * taux / 100)}\n"
          f"Rubrique URSSAF : « Prestations de services BNC ». Déclarer même si zéro.")


def relances(a):
    rows = lire()
    today = dt.date.today()
    due = [r for r in rows if not r["encaisse_le"] and dt.date.fromisoformat(r["echeance"]) < today]
    if not due:
        print("Aucune facture en retard.")
        return
    for r in due:
        retard = (today - dt.date.fromisoformat(r["echeance"])).days
        print(f"--- {r['numero']} · {r['client']} · {eur(r['montant'])} · {retard} jours de retard\n"
              f"Objet : Facture {r['numero']} échue le {dt.date.fromisoformat(r['echeance']):%d/%m/%Y}\n\n"
              f"Bonjour,\n\nSauf erreur de ma part, la facture {r['numero']} d'un montant de {eur(r['montant'])}, "
              f"relative à « {r['objet']} », reste impayée à ce jour. Je vous remercie de procéder au règlement "
              f"sous 8 jours. Si le paiement est déjà parti, merci d'ignorer ce message.\n\nCordialement,\n{cfg()['nom']}\n")


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    n = sub.add_parser("nouvelle")
    n.add_argument("--client", required=True); n.add_argument("--adresse", required=True)
    n.add_argument("--objet", required=True); n.add_argument("--montant", required=True, type=float)
    n.add_argument("--siret-client", default=""); n.add_argument("--detail", default="")
    n.set_defaults(f=nouvelle)
    e = sub.add_parser("encaissee"); e.add_argument("numero"); e.add_argument("--date", default="")
    e.add_argument("--mode", default="virement"); e.set_defaults(f=encaissee)
    u = sub.add_parser("urssaf"); u.add_argument("--trimestre", default=""); u.set_defaults(f=urssaf)
    r = sub.add_parser("relances"); r.set_defaults(f=relances)
    a = ap.parse_args()
    a.f(a)


if __name__ == "__main__":
    main()
